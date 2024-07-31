import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Side

def create_test_result_file(test_scenario_df, tla, env, ROOT_DIR, usernames, source_conn, target_conn, source_host, target_host, tre_userlist):
    TEMP_FILE_NAME = 'C:\\SVAU VI Techrem\\automation_sv_setup\\2019_templates\\2019 TLA ENV Test Report Structural Validation.xlsx'
    print("TEMP FILE NAME line 441:", TEMP_FILE_NAME)

    test_file = "test file tla env Test Report Structural Validation.xlsx"
    test_file_loc = os.path.join(ROOT_DIR, test_file)

    try:
        # Load workbook
        wb = load_workbook(filename=TEMP_FILE_NAME)

        # Creating border format
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Sheet: Test Summary [0]
        ws0 = wb.worksheets[0]

        # Read data from TEMP_FILE_NAME
        ts_df = pd.read_excel(TEMP_FILE_NAME, sheet_name="Test Summary")

        # Replace placeholders in DataFrame
        ts_df['tc_name'] = ts_df['tc_name'].str.replace("<TLA>", tla)
        ts_df['tc_name'] = ts_df['tc_name'].str.replace("<ENV>", env)

        # Write data to the Excel sheet
        for i in range(len(ts_df)):
            cell = "A" + str(i + 2)
            ws0[cell] = ts_df.iloc[i]['tc_name']
            # Apply border format
            ws0[cell].border = border

        # Save the workbook to the new location
        wb.save(test_file_loc)
        print(f"File saved to {test_file_loc}")

    except Exception as e:
        print(f"An error occurred: {e}")
